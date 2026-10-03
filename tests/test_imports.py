from decimal import Decimal
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from core.exceptions import BusinessRuleError
from ledger.models import Ledger, Voucher
from masters.models import SKU, Colour, Material, Party, Style
from masters.services import imports, parties
from tests.conftest import make_user

D = Decimal


def sheet(kind, *rows, header=None):
    wb = Workbook()
    ws = wb.active
    ws.append(header or imports.TEMPLATES[kind]["columns"])
    for r in rows:
        ws.append(list(r))
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def run(kind, data, company, user, **kw):
    return imports.run_import(kind=kind, fileobj=data, company=company, user=user, **kw)


def test_templates_download_as_workbooks_with_the_right_headers():
    for kind, spec in imports.TEMPLATES.items():
        ws = load_workbook(BytesIO(imports.build_template(kind))).worksheets[0]
        assert [c.value for c in ws[1]] == spec["columns"]


def test_party_import_creates_ledgers_and_numeric_mobiles_work(company, owner):
    data = sheet("parties",
                 ["Mehta Traders", "customer", 9876543210, "Rakesh", "", "03ABCDE1234F1Z5", "", "wholesaler", 50000, 30, 0],
                 ["Sharma Stitching", "fabricator, vendor", "9811111111", "", "", "", "", "", "", "", ""])
    r = run("parties", data, company, owner, commit=True)
    assert r.clean and r.committed and r.ok == 2
    mehta = Party.objects.get(name="Mehta Traders")
    assert mehta.mobile == "9876543210" and mehta.credit_limit == D("50000") and mehta.category == "wholesaler"
    assert Party.objects.get(name="Sharma Stitching").payable_ledger is not None


def test_check_only_saves_nothing(company, owner):
    data = sheet("parties", ["Mehta Traders", "customer", "9876543210", "", "", "", "", "", "", "", ""])
    r = run("parties", data, company, owner, commit=False)
    assert r.clean and not r.committed and r.ok == 1
    assert not Party.objects.exists() and not Ledger.objects.filter(name="Mehta Traders").exists()


def test_any_bad_row_means_nothing_is_imported_and_errors_name_the_rows(company, owner):
    data = sheet("parties",
                 ["Good Firm", "customer", "9876543210", "", "", "", "", "", "", "", ""],
                 ["Dup Mobile", "customer", "9876543210", "", "", "", "", "", "", "", ""],
                 ["Bad Role", "boss", "9811111111", "", "", "", "", "", "", "", ""],
                 ["Bad GSTIN", "customer", "9822222222", "", "", "XYZ", "", "", "", "", ""])
    r = run("parties", data, company, owner, commit=True)
    assert not r.committed and r.ok == 1
    assert dict(r.errors)[3].startswith("Mobile 9876543210 already belongs to customer Good Firm")
    assert "Unknown role" in dict(r.errors)[4] and 5 in dict(r.errors)
    assert not Party.objects.exists()  # the good row was rolled back too


def test_style_import_makes_skus_and_new_colours(company, owner):
    data = sheet("styles", ["jgr-200", "Slim jogger", "JGR", "Black, Olive Green", "S, M, L", "6103", 699, "Slim fit"])
    r = run("styles", data, company, owner, commit=True)
    assert r.clean
    style = Style.objects.get(style_no="JGR-200")
    assert SKU.objects.filter(style=style).count() == 6 and Colour.objects.filter(name="Olive Green").exists()
    assert style.mrp == D("699") and style.hsn.code == "6103"


def test_style_import_reports_unknown_size_product_and_duplicates(company, owner):
    data = sheet("styles",
                 ["A1", "a", "TRK", "Black", "S", "", "", ""],
                 ["A1", "dup", "TRK", "Black", "S", "", "", ""],
                 ["A2", "a", "NOPE", "Black", "S", "", "", ""],
                 ["A3", "a", "TRK", "Black", "XXXL", "", "", ""],
                 ["A4", "a", "TRK", "Black", "S", "9999", "", ""])
    r = run("styles", data, company, owner, commit=True)
    errs = dict(r.errors)
    assert set(errs) == {3, 4, 5, 6} and "Unknown product" in errs[4] and "Unknown size" in errs[5] and "Unknown HSN" in errs[6]
    assert not Style.objects.exists()


def test_material_import(company, owner):
    r = run("materials", sheet("materials", ["FAB-1", "Fleece", "fabric", "kg", "cotton", 280, 180],
                               ["ZIP-1", "Zipper", "trim", "PCS", "", "", ""]), company, owner, commit=True)
    assert r.clean and Material.objects.get(code="FAB-1").gsm == 280 and Material.objects.get(code="ZIP-1").unit.code == "PCS"
    bad = run("materials", sheet("materials", ["X", "x", "cloth", "KG", "", "", ""]), company, owner, commit=True)
    assert "kind must be" in dict(bad.errors)[2]


def test_opening_balance_import_posts_one_balanced_voucher_with_bill_refs(company, factory, owner):
    parties.create_party(company=company, name="Mehta Traders", mobile="9876543210", is_customer=True)
    data = sheet("opening", ["Cash", 5000, "", "", ""], ["Mehta Traders", 20000, "", "OLD-7", "2026-05-15"],
                 ["Profit & Loss A/c", "", 20000, "", ""])
    r = run("opening", data, company, owner, factory=factory, commit=True)
    assert r.clean and r.committed
    v = Voucher.objects.get(voucher_type="opening")
    assert v.lines.count() == 4 and v.status == "posted"  # includes the Opening Balance Difference line
    from ledger.selectors import outstanding_bills

    assert outstanding_bills(Party.objects.get().customer_ledger)["bills"] == {"OLD-7": D("20000.00")}


def test_opening_import_errors_and_missing_factory(company, factory, owner):
    r = run("opening", sheet("opening", ["No such ledger", 10, "", "", ""], ["Cash", 5, 5, "", ""]), company, owner,
            factory=factory, commit=True)
    assert set(dict(r.errors)) == {2, 3} and not Voucher.objects.exists()
    with pytest.raises(BusinessRuleError):
        run("opening", sheet("opening", ["Cash", 10, "", "", ""]), company, owner, commit=True)


def test_wrong_headers_and_non_excel_files_are_rejected_clearly(company, owner):
    with pytest.raises(BusinessRuleError, match="missing"):
        run("parties", sheet("parties", header=["name", "mobile"]), company, owner)
    with pytest.raises(BusinessRuleError, match="could not be read"):
        run("parties", BytesIO(b"not an excel file"), company, owner)


def test_import_screen_flow_and_permissions(company, factory, owner, accountant):
    c = Client()
    c.force_login(accountant)
    assert c.get(reverse("excel_import")).status_code == 200
    assert c.get(reverse("import_template", args=["parties"])).status_code == 200
    assert c.get(reverse("import_template", args=["styles"])).status_code == 404  # accountant cannot create styles
    data = sheet("parties", ["Mehta Traders", "customer", "9876543210", "", "", "", "", "", "", "", ""])
    upload = SimpleUploadedFile("p.xlsx", data.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r = c.post(reverse("excel_import"), {"kind": "parties", "mode": "check", "file": upload})
    assert b"Check passed" in r.content and not Party.objects.exists()
    data.seek(0)
    upload = SimpleUploadedFile("p.xlsx", data.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r = c.post(reverse("excel_import"), {"kind": "parties", "mode": "import", "file": upload})
    assert b"Imported 1 rows" in r.content and Party.objects.count() == 1
    r = c.post(reverse("excel_import"), {"kind": "styles", "mode": "import", "file": upload})
    assert r.status_code == 403
    stranger = make_user("nobody")
    s = Client()
    s.force_login(stranger)
    assert s.get(reverse("excel_import")).status_code == 403
