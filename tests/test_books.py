"""The books as reports (E9.5, ACC-04/06/07/13/19): ledger statement, day book, P&L, balance sheet, ageing."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from core.models import FinancialYear, Role
from ledger.models import Ledger
from ledger.services.posting import AllocationSpec, LineSpec, post_voucher
from reports.services import books
from tests import sales_helpers as sh
from tests.conftest import make_user
from tests.test_sales import quick_invoice

D = Decimal
DAY = sh.DAY
FY_FROM, FY_TO = date(2026, 4, 1), date(2027, 3, 31)


@pytest.fixture
def ns(company, factory, owner):
    return sh.build(company, factory, owner)


def led(company, key):
    return Ledger.objects.get(company=company, system_key=key)


@pytest.fixture
def month(ns, company, factory, owner):
    """One sale of 5,000 costing 3,000, 500 of office expense paid in cash, and 2,000 received against the invoice."""
    inv = quick_invoice(ns, qty="10", rate="500")
    office = led(company, "office_expenses")
    post_voucher(company=company, factory=factory, voucher_type="journal", date=DAY, user=owner, lines=[
        LineSpec(ledger=office, debit=D("500.00")), LineSpec(ledger=led(company, "cash"), credit=D("500.00"))])
    post_voucher(company=company, factory=factory, voucher_type="receipt", date=DAY + timedelta(days=2), user=owner, lines=[
        LineSpec(ledger=led(company, "cash"), debit=D("2000.00")),
        LineSpec(ledger=ns.local.customer_ledger, credit=D("2000.00"), allocations=(AllocationSpec("against", D("2000.00"), inv.number),))])
    return inv


def pl(company, owner, factory=None, f=FY_FROM, t=FY_TO):
    return books.profit_and_loss(company, user=owner, factory=factory, date_from=f, date_to=t)


def test_profit_and_loss_gross_to_net_per_factory_and_per_period(month, company, factory, factory2, owner):
    r = pl(company, owner)
    assert [(x["ledger"], x["amount"]) for x in r["income"]] == [("Sales - Ready Stock", D("5000.00"))]
    assert [(x["ledger"], x["amount"]) for x in r["expense"]] == [("Cost of Goods Sold", D("3000.00"))]
    assert (r["gross_profit"], r["net_profit"], r["gross_margin"], r["net_margin"]) == (D("2000.00"), D("1500.00"), D("40.0"), D("30.0"))
    assert [(x["ledger"], x["amount"]) for x in r["other_expense"]] == [("Office Expenses", D("500.00"))]
    assert pl(company, owner, factory)["net_profit"] == D("1500.00")
    assert pl(company, owner, factory2)["net_profit"] == D("0.00") and pl(company, owner, factory2)["income"] == []
    assert pl(company, owner, f=DAY + timedelta(days=1))["net_profit"] == D("0.00")          # nothing after the sale
    assert pl(company, owner, f=FY_FROM, t=DAY - timedelta(days=1))["net_profit"] == D("0.00")


def test_comparison_with_the_same_days_last_year(month, company, owner):
    c = books.compare_profit_and_loss(company, user=owner, date_from=FY_FROM, date_to=FY_TO)
    assert c["now"]["net_profit"] == D("1500.00") and c["before"]["net_profit"] == D("0.00")
    assert (c["before_from"], c["before_to"]) == (date(2025, 4, 1), date(2026, 3, 31))
    sales = next(b for b in c["lines"] if b["section"].startswith("Sales"))
    assert sales["rows"] == [["Sales - Ready Stock", D("5000.00"), D("0.00")]]
    assert books.previous_year_range(date(2024, 2, 29), date(2024, 3, 31)) == (date(2023, 2, 28), date(2023, 3, 31))


def test_the_balance_sheet_balances_and_splits_profit_by_financial_year(month, company, factory, factory2, owner):
    bs = books.balance_sheet(company, user=owner, as_of=DAY + timedelta(days=5))
    assert bs["tallies"] and bs["current_profit"] == D("1500.00") and bs["prior_profit"] == D("0.00")
    assets = {a["ledger"]: a["amount"] for a in bs["assets"]}
    assert assets["Finished Goods Stock"] == D("237000.00") and assets["Cash"] == D("1500.00")
    assert assets["Punjab Traders"] == D("3000.00")
    assert bs["assets_total"] == bs["liabilities_total"] == D("241500.00")
    FinancialYear.objects.create(company=company, start_date=date(2027, 4, 1), end_date=date(2028, 3, 31), label="27-28")
    nxt = books.balance_sheet(company, user=owner, as_of=date(2027, 5, 1))                  # a year later: it is now earlier-year profit
    assert nxt["tallies"] and nxt["prior_profit"] == D("1500.00") and nxt["current_profit"] == D("0.00")
    assert books.balance_sheet(company, user=owner, factory=factory2, as_of=DAY)["assets"] == []
    assert books.balance_sheet(company, user=owner, factory=factory, as_of=DAY)["tallies"]


def test_ledger_statement_runs_a_balance_and_carries_the_opening(month, ns, company, factory, owner):
    debtors = ns.local.customer_ledger
    st = books.ledger_statement(debtors, user=owner, date_from=FY_FROM, date_to=FY_TO)
    assert st["opening"] == 0 and [(r["debit"], r["credit"], r["balance"]) for r in st["rows"]] == [
        (D("5000.00"), D("0.00"), D("5000.00")), (D("0.00"), D("2000.00"), D("3000.00"))]
    assert st["closing"] == D("3000.00") and "Sales - Ready Stock" in st["rows"][0]["particulars"]
    later = books.ledger_statement(debtors, user=owner, date_from=DAY + timedelta(days=1), date_to=FY_TO)
    assert later["opening"] == D("5000.00") and len(later["rows"]) == 1 and later["closing"] == D("3000.00")
    assert books.ledger_statement(debtors, user=owner, factory=None, date_from=FY_FROM, date_to=DAY)["closing"] == D("5000.00")


def test_day_book_lists_posted_vouchers_by_date_and_type(month, company, owner):
    book = books.day_book(company, user=owner, date_from=DAY, date_to=DAY)
    assert sorted(v.voucher_type for v in book["vouchers"]) == ["journal", "sales"]
    assert [v.voucher_type for v in books.day_book(company, user=owner, date_from=DAY, date_to=FY_TO, voucher_type="receipt")["vouchers"]] == ["receipt"]


def test_ageing_buckets_by_days_past_due_and_shows_advances_apart(month, ns, company, factory, owner):
    due = DAY + timedelta(days=30)
    ag = books.ageing(company, user=owner, kind="debtors", as_of=DAY + timedelta(days=10))
    assert ag["total"] == D("3000.00") and ag["totals"]["Not due"] == D("3000.00")
    ag = books.ageing(company, user=owner, kind="debtors", as_of=due + timedelta(days=10))
    assert ag["totals"]["1-30"] == D("3000.00") and ag["parties"][0]["bills"][0]["days"] == 10
    ag = books.ageing(company, user=owner, kind="debtors", as_of=due + timedelta(days=100))
    assert ag["totals"]["91-180"] == D("3000.00")
    post_voucher(company=company, factory=factory, voucher_type="receipt", date=DAY, user=owner, lines=[
        LineSpec(ledger=led(company, "cash"), debit=D("700.00")),
        LineSpec(ledger=ns.far.customer_ledger, credit=D("700.00"), allocations=(AllocationSpec("advance", D("700.00")),))])
    ag = books.ageing(company, user=owner, kind="debtors", as_of=DAY)
    assert ag["unadjusted"] == D("-700.00") and ag["total"] == D("5000.00")        # as at DAY: the 2,000 comes two days later
    assert books.ageing(company, user=owner, kind="creditors", as_of=DAY)["parties"] == []


def test_every_report_opens_exports_and_respects_factories(month, ns, company, factory, factory2, owner):
    c = Client()
    c.force_login(owner)
    pages = [("trial_balance", []), ("ledger_pick", []), ("ledger_statement", [ns.local.customer_ledger.pk]), ("day_book", []),
             ("profit_loss", []), ("balance_sheet", []), ("ageing", ["debtors"]), ("ageing", ["creditors"])]
    for name, args in pages:
        r = c.get(reverse(name, args=args))
        assert r.status_code == 200, name
    assert c.get(reverse("profit_loss"), {"compare": "1"}).status_code == 200
    assert c.get(reverse("day_book"), {"type": "journal"}).status_code == 200
    assert "Does not" not in c.get(reverse("balance_sheet")).content.decode()
    page = c.get(reverse("trial_balance")).content.decode()
    assert reverse("ledger_statement", args=[ns.local.customer_ledger.pk]) in page          # drill from the trial balance
    st = c.get(reverse("ledger_statement", args=[ns.local.customer_ledger.pk])).content.decode()
    assert "Opening balance" in st and "5000.00" in st and "3000.00" in st
    bad = c.get(reverse("profit_loss"), {"from": "01/04/2026"})
    assert bad.status_code == 200 and "is not a date" in bad.content.decode()
    for name, args, title in [("profit_loss", [], "Profit and loss"), ("balance_sheet", [], "Balance sheet"), ("ageing", ["debtors"], "Receivables"),
                              ("day_book", [], "Day book"), ("ledger_statement", [ns.local.customer_ledger.pk], "Punjab Traders")]:
        r = c.get(reverse(name, args=args), {"format": "xlsx"})
        assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"], name
        ws = load_workbook(BytesIO(r.content)).active
        assert ws["A1"].value == title
    ws = load_workbook(BytesIO(c.get(reverse("profit_loss"), {"format": "xlsx"}).content)).active
    cells = {row[0].value: row[1].value for row in ws.iter_rows(min_row=3) if row[0].value}
    assert cells["Net profit"] == D("1500.00") and cells["Gross profit"] == D("2000.00")
    # factory 2 sees nothing of factory 1's books
    other = make_user("books2")
    other.roles.add(Role.objects.get(name="Accountant"))
    other.allowed_factories.add(factory2)
    oc = Client()
    oc.force_login(other)
    assert "Punjab Traders" not in oc.get(reverse("ageing", args=["debtors"])).content.decode()
    assert "5000.00" not in oc.get(reverse("profit_loss")).content.decode()
    clerk = make_user("clerk3")
    clerk.roles.add(Role.objects.get(name="Billing Clerk"))
    clerk.allowed_factories.add(factory)
    cc = Client()
    cc.force_login(clerk)
    assert cc.get(reverse("profit_loss")).status_code == 403 and cc.get(reverse("balance_sheet")).status_code == 403


def test_ledger_book_shows_each_chosen_ledger_with_combined_totals(month, ns, company, owner):
    debtors, cash = ns.local.customer_ledger, led(company, "cash")
    book = books.ledger_book([debtors, cash], user=owner, date_from=FY_FROM, date_to=FY_TO)
    assert [s["ledger"] for s in book["statements"]] == [debtors, cash]
    assert book["debit"] == sum((s["debit"] for s in book["statements"]), D("0.00"))
    assert book["closing"] == D("3000.00") + D("1500.00")                    # debtors 3,000 Dr, cash 1,500 Dr
    c = Client()
    c.force_login(owner)
    page = c.get(reverse("ledger_book"), {"ledger": [debtors.pk, cash.pk, 999999], "from": "2026-04-01", "to": "2027-03-31"}).content.decode()
    assert "Punjab Traders" in page and "Cash" in page and "All selected" in page and "2 selected" in page
    assert c.get(reverse("ledger_book")).status_code == 200                  # no ledger chosen yet: just the picker
    r = c.get(reverse("ledger_book"), {"ledger": [debtors.pk, cash.pk], "format": "xlsx"})
    assert "spreadsheetml" in r["Content-Type"] and load_workbook(BytesIO(r.content)).active["A1"].value == "Ledger book"
    assert c.get(reverse("ledger_book"), {"format": "xlsx"}).status_code == 302
