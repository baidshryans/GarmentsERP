"""Excel import for migration (PRD 8.1): parties, styles, materials and opening balances.

Every import is all-or-nothing: rows are validated by the same services the screens use, inside one
transaction. "Check only" runs everything and rolls back, so users can fix the sheet before importing.
"""
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook

from core.exceptions import BusinessRuleError
from ledger.models import Ledger
from ledger.services.opening import OpeningEntry, post_opening_balances
from masters.models import Colour, Material, Product, Size, Unit
from masters.services import parties, styles
from tax.models import HSN

TEMPLATES = {
    "parties": {
        "title": "Parties",
        "columns": ["name", "roles", "mobile", "contact_person", "email", "gstin", "pan", "category",
                    "credit_limit", "credit_days", "discount_pct"],
        "example": ["Mehta Traders", "customer", "9876543210", "Rakesh Mehta", "", "03ABCDE1234F1Z5", "", "wholesaler", 50000, 30, 0],
        "notes": "roles: customer, vendor, fabricator, agent, transporter (comma separated). category: wholesaler, "
                 "distributor, retailer, online, walkin, institutional. A duplicate customer mobile is rejected.",
    },
    "styles": {
        "title": "Styles",
        "columns": ["style_no", "name", "product", "colours", "sizes", "hsn", "mrp", "description"],
        "example": ["JGR-104", "Cuffed jogger", "JGR", "Black, Navy, Grey Melange", "S, M, L, XL, XXL", "6103", 699, ""],
        "notes": "product: an existing product code (TRK, JGR, TSH, SET). colours: comma separated, new ones are created. "
                 "sizes: existing size codes. Every colour-size pair gets a SKU and barcode.",
    },
    "materials": {
        "title": "Materials",
        "columns": ["code", "name", "kind", "unit", "composition", "gsm", "width_cm"],
        "example": ["FAB-001", "Cotton fleece 280 GSM", "fabric", "KG", "80% cotton 20% polyester", 280, 180],
        "notes": "kind: fabric, trim or packing. unit: an existing unit code (KG, MTR, PCS, GRS ...).",
    },
    "opening": {
        "title": "Opening balances",
        "columns": ["ledger", "debit", "credit", "reference", "due_date"],
        "example": ["Cash", 5000, "", "", ""],
        "notes": "ledger: the exact ledger name. Fill debit or credit. For debtors and creditors give a bill reference "
                 "and optional due date (YYYY-MM-DD). Any difference goes to Opening Balance Difference.",
    },
}


@dataclass
class ImportResult:
    kind: str
    rows: int = 0
    ok: int = 0
    errors: list = field(default_factory=list)  # (row number, message)
    committed: bool = False

    @property
    def clean(self):
        return not self.errors


def build_template(kind: str) -> bytes:
    spec = TEMPLATES[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = spec["title"]
    ws.append(spec["columns"])
    ws.append(spec["example"])
    notes = wb.create_sheet("Notes")
    notes.append(["How to fill this sheet"])
    notes.append([spec["notes"]])
    notes.append(["Row 2 is an example; delete or overwrite it. Keep the header row exactly as it is."])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 4)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        value = value.date()
    return str(value).strip()


def _read(fileobj, kind):
    columns = TEMPLATES[kind]["columns"]
    try:
        ws = load_workbook(fileobj, read_only=True, data_only=True).worksheets[0]
    except Exception:
        raise BusinessRuleError("That file could not be read. Upload an .xlsx file made from the template.")
    rows = ws.iter_rows(values_only=True)
    header = [_text(h).lower() for h in next(rows, [])]
    missing = [c for c in columns if c not in header]
    if missing:
        raise BusinessRuleError("The header row is missing: " + ", ".join(missing) + ". Download the template again.")
    idx = {c: header.index(c) for c in columns}
    out = []
    for n, row in enumerate(rows, start=2):
        values = {c: (row[idx[c]] if idx[c] < len(row) else None) for c in columns}
        if all(_text(v) == "" for v in values.values()):
            continue
        out.append((n, values))
    return out


def _dec(value, label, default=None):
    t = _text(value).replace(",", "")
    if not t:
        if default is not None:
            return default
        raise ValueError(f"{label} is required.")
    try:
        return Decimal(t)
    except InvalidOperation:
        raise ValueError(f"{label} '{t}' is not a number.")


def _split(value):
    return [p.strip() for p in _text(value).replace(";", ",").split(",") if p.strip()]


CATEGORY = {"wholesaler": "wholesaler", "distributor": "distributor", "retailer": "retailer", "online": "online",
            "online seller": "online", "walkin": "walkin", "walk-in": "walkin", "institutional": "institutional"}
ROLE = {"customer": "is_customer", "vendor": "is_vendor", "fabricator": "is_fabricator",
        "agent": "is_agent", "transporter": "is_transporter"}


def _row_parties(company, v, **ctx):
    roles = {}
    for r in _split(v["roles"]):
        if r.lower() not in ROLE:
            raise ValueError(f"Unknown role '{r}'.")
        roles[ROLE[r.lower()]] = True
    cat = _text(v["category"]).lower()
    if cat and cat not in CATEGORY:
        raise ValueError(f"Unknown category '{cat}'.")
    parties.create_party(
        company=company, name=_text(v["name"]), mobile=_text(v["mobile"]), contact_person=_text(v["contact_person"]),
        email=_text(v["email"]), gstin=_text(v["gstin"]), pan=_text(v["pan"]), category=CATEGORY.get(cat, ""),
        credit_limit=_dec(v["credit_limit"], "credit_limit", Decimal("0")),
        credit_days=int(_dec(v["credit_days"], "credit_days", Decimal("0"))),
        discount_pct=_dec(v["discount_pct"], "discount_pct", Decimal("0")), **roles,
    )


def _row_styles(company, v, **ctx):
    product = Product.objects.filter(code__iexact=_text(v["product"])).first()
    if product is None:
        raise ValueError(f"Unknown product '{_text(v['product'])}'.")
    size_objs = []
    for code in _split(v["sizes"]):
        size = Size.objects.filter(code__iexact=code).first()
        if size is None:
            raise ValueError(f"Unknown size '{code}'. Add it under Sizes first.")
        size_objs.append(size)
    colour_objs = [Colour.objects.filter(name__iexact=n).first() or Colour.objects.create(name=n)
                   for n in _split(v["colours"])]
    hsn = None
    if _text(v["hsn"]):
        hsn = HSN.objects.filter(code=_text(v["hsn"])).first()
        if hsn is None:
            raise ValueError(f"Unknown HSN '{_text(v['hsn'])}'.")
    mrp = _dec(v["mrp"], "mrp", Decimal("0")) or None
    styles.create_style(
        company=company, style_no=_text(v["style_no"]), product=product, name=_text(v["name"]),
        colours=colour_objs, sizes=size_objs, hsn=hsn, mrp=mrp, description=_text(v["description"]),
    )


def _row_materials(company, v, **ctx):
    unit = Unit.objects.filter(code__iexact=_text(v["unit"])).first()
    if unit is None:
        raise ValueError(f"Unknown unit '{_text(v['unit'])}'.")
    kind = _text(v["kind"]).lower()
    if kind not in {k for k, _ in Material.Kind.choices}:
        raise ValueError(f"kind must be fabric, trim or packing (got '{kind}').")
    gsm = _text(v["gsm"])
    m = Material(
        code=_text(v["code"]), name=_text(v["name"]), kind=kind, unit=unit, composition=_text(v["composition"]),
        gsm=int(_dec(v["gsm"], "gsm")) if gsm else None, width_cm=_dec(v["width_cm"], "width_cm") if _text(v["width_cm"]) else None,
    )
    m.full_clean()
    m.save()


ROW_HANDLERS = {"parties": _row_parties, "styles": _row_styles, "materials": _row_materials}


def run_import(*, kind, fileobj, company, user, factory=None, commit=False) -> ImportResult:
    """Validate (and, with commit=True and no errors, save) every row in one transaction."""
    result = ImportResult(kind=kind)
    rows = _read(fileobj, kind)
    result.rows = len(rows)
    with transaction.atomic():
        if kind == "opening":
            _opening(rows, result, company, user, factory)
        else:
            handler = ROW_HANDLERS[kind]
            for n, values in rows:
                try:
                    with transaction.atomic():
                        handler(company, values)
                    result.ok += 1
                except (ValueError, BusinessRuleError) as exc:
                    result.errors.append((n, str(exc)))
                except ValidationError as exc:
                    result.errors.append((n, "; ".join(exc.messages)))
        if result.errors or not commit:
            transaction.set_rollback(True)
        else:
            result.committed = True
    return result


def _opening(rows, result, company, user, factory):
    if factory is None:
        raise BusinessRuleError("Choose the factory these opening balances belong to.")
    entries = []
    ledgers = {l.name.lower(): l for l in Ledger.objects.filter(company=company, is_active=True)}
    for n, v in rows:
        try:
            ledger = ledgers.get(_text(v["ledger"]).lower())
            if ledger is None:
                raise ValueError(f"No active ledger named '{_text(v['ledger'])}'.")
            dr, cr = _dec(v["debit"], "debit", Decimal("0")), _dec(v["credit"], "credit", Decimal("0"))
            if bool(dr) == bool(cr):
                raise ValueError("Fill either debit or credit.")
            due = _text(v["due_date"])
            entries.append(OpeningEntry(
                ledger=ledger, debit=dr, credit=cr, reference=_text(v["reference"]),
                due_date=date.fromisoformat(due) if due else None,
            ))
            result.ok += 1
        except ValueError as exc:
            result.errors.append((n, str(exc)))
    if not result.errors:
        try:
            with transaction.atomic():
                post_opening_balances(company=company, factory=factory, entries=entries, user=user)
        except (BusinessRuleError, ValidationError) as exc:
            result.errors.append((0, "; ".join(exc.messages) if isinstance(exc, ValidationError) else str(exc)))
