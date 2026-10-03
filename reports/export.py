"""Excel export for reports (E10.2). PDF is the browser's Print dialog: report pages print plain black on white."""
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_response(filename, title, headers, rows, notes=()):
    """One sheet: a title line, optional notes, a header row, then the rows. Amounts stay numbers so Excel can add them."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    for n in notes:
        ws.append([n])
    ws.append([])
    ws.append(list(headers))
    head = ws.max_row
    for c in ws[head]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([v if isinstance(v, (int, float, Decimal, str, type(None))) or hasattr(v, "isoformat") else str(v) for v in row])
    for col in ws.iter_cols(min_row=head, max_row=ws.max_row):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 50)
        for c in col[1:]:
            if isinstance(c.value, Decimal):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
    resp = HttpResponse(content_type=XLSX)
    resp["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    wb.save(resp)
    return resp
